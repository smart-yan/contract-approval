"""测试用的最小文档工厂。

**不依赖仓库外的样例文件** —— 每个用例自己现场生成，所以测试在任何机器上都能跑，
也不会因为样例文件被改动而悄悄失效。
"""

from __future__ import annotations

import io
from collections.abc import Iterable
from pathlib import Path

import docx

from app.parsers.docx_parser import normalize_text
from app.schemas.document import Paragraph, ParseResult

#: 一个 block 可以是：
#:   ``str``             → 一个普通段落
#:   ``tuple[str, ...]`` → **一个段落**，各部分之间是作者插入的软换行（``<w:br/>``）
#:   ``list[list[str]]`` → 一张表（外层行、内层单元格）
Block = str | tuple[str, ...] | list[list[str]]


def soft_break_paragraph(*parts: str) -> tuple[str, ...]:
    """构造"一个段落内含软换行"的输入。

    这种段落是 DOCX 位置契约最容易踩的坑：python-docx 会把 ``<w:br/>`` 读成 ``\\n``，
    于是段落文本自带换行符，段落边界就再也分不出来了。
    """
    return tuple(parts)


def save_docx(path: Path, *blocks: Block) -> Path:
    """按给定顺序生成一个真实 DOCX 并写到 ``path``，返回该路径。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    _build(*blocks).save(str(path))
    return path


def docx_bytes(*blocks: Block) -> bytes:
    """同 :func:`save_docx`，但返回字节流（用于构造 multipart 上传体）。"""
    buffer = io.BytesIO()
    _build(*blocks).save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
def _build(*blocks: Block) -> docx.Document:
    document = docx.Document()
    for block in blocks:
        if isinstance(block, str):
            document.add_paragraph(block)
        elif isinstance(block, tuple):
            _add_soft_break_paragraph(document, block)
        else:
            _add_table(document, block)
    return document


def _add_soft_break_paragraph(document: docx.Document, parts: tuple[str, ...]) -> None:
    """一个段落，各部分之间插入 ``<w:br/>``（读回来就是段落内的 ``\\n``）。"""
    paragraph = document.add_paragraph()
    for i, part in enumerate(parts):
        if i:
            paragraph.add_run().add_break()
        paragraph.add_run(part)


def _add_table(document: docx.Document, rows: Iterable[Iterable[str]]) -> None:
    rows = [list(row) for row in rows]
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            table.cell(r, c).text = cell_text


# --------------------------------------------------------------------------- #
# 直接构造 ParseResult（不经过 DOCX）
#
# 理解层的纯函数测试不需要走真实的 DOCX 编解码 —— 直接给出段落序列，
# 又快又能把"输入是什么"写得一目了然。序号由 make_parse_result 统一分配，
# 避免手写下标写错。
# --------------------------------------------------------------------------- #
def para(text: str) -> Paragraph:
    """一个正文段落（``PARAGRAPH``）。

    ⚠️ 文本会经过与真实解析器**同一套**清洗（``normalize_text``）。
    这不是多此一举：``ParseResult`` 的契约规定段落文本已压成单行，
    工厂若放行 ``"   "`` 这类原文，就会造出**生产环境不可能出现**的输入，
    让测试断言到不存在的行为。
    """
    return Paragraph(index=-1, block_type="PARAGRAPH", text=normalize_text(text))


def row(*cells: str) -> Paragraph:
    """一个表格行（``TABLE_ROW``），单元格各自清洗后用制表符连接 —— 与解析器一致。"""
    return Paragraph(
        index=-1,
        block_type="TABLE_ROW",
        text="\t".join(normalize_text(cell) for cell in cells),
    )


def make_parse_result(*blocks: Paragraph) -> ParseResult:
    """把若干段落编成一个 ``ParseResult``，序号按顺序重排。

    ``status`` / ``text`` 的算法与真实解析器一致 —— 测试夹具不引入第二套规则。
    """
    paragraphs = [b.model_copy(update={"index": i}) for i, b in enumerate(blocks)]
    text = "\n".join(p.text for p in paragraphs)
    return ParseResult(
        status="PARSED" if text.strip() else "EMPTY",
        parser="DocxParser",
        source_file_type="DOCX",
        text=text,
        paragraphs=paragraphs,
    )


__all__ = [
    "Block",
    "docx_bytes",
    "make_parse_result",
    "para",
    "row",
    "save_docx",
    "soft_break_paragraph",
]
